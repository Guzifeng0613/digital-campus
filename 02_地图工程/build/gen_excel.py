# -*- coding: utf-8 -*-
"""生成建筑点位台账.xlsx（含建筑、自然/POI、分类统计三个表）。"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import campus_data as C

recs = C.numbered_buildings()
cat_zh = {k: v[0] for k, v in C.CATEGORIES.items()}
cat_color = {k: v[1] for k, v in C.CATEGORIES.items()}

wb = Workbook()
thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="2E4053")
head_font = Font(color="FFFFFF", bold=True, size=11)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ---- Sheet1 建筑 ----
ws = wb.active; ws.title = "建筑点位台账"
head = ["编号", "ID", "建筑名称", "类别", "状态", "网格X", "网格Y", "宽(瓦片)", "深(瓦片)", "层数", "形状", "功能说明"]
ws.append(head)
for c in range(1, len(head)+1):
    cell = ws.cell(1, c); cell.fill = head_fill; cell.font = head_font; cell.alignment = center; cell.border = border
for b in recs:
    gx, gy, gw, gh = b["grid"]
    row = ["%02d" % b["no"], b["id"], b["name"], cat_zh[b["cat"]],
           {"built": "已建", "building": "在建"}.get(b["status"], b["status"]),
           gx, gy, gw, gh, b.get("floors", 0), b["shape"], b["desc"]]
    ws.append(row)
    r = ws.max_row
    for c in range(1, len(head)+1):
        cell = ws.cell(r, c); cell.border = border
        cell.alignment = left if c == 12 else center
    # 分类色
    color = cat_color[b["cat"]].lstrip("#")
    ws.cell(r, 4).fill = PatternFill("solid", fgColor=color)
    if b["status"] == "building":
        ws.cell(r, 5).fill = PatternFill("solid", fgColor="FAD7A0")
widths = [6, 16, 20, 12, 7, 8, 8, 9, 9, 6, 9, 50]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 22

# ---- Sheet2 自然/POI ----
ws2 = wb.create_sheet("自然与兴趣点")
h2 = ["类型", "名称", "类别/类型", "网格X", "网格Y"]
ws2.append(h2)
for c in range(1, len(h2)+1):
    cell = ws2.cell(1, c); cell.fill = head_fill; cell.font = head_font; cell.alignment = center; cell.border = border
for n in C.build_nature():
    gx, gy, *_ = n["grid"]
    ws2.append(["自然/区域", n["name"], cat_zh[n["cat"]], gx, gy])
for p in C.build_pois():
    ws2.append(["兴趣点", p["name"], p["type"], p["grid"][0], p["grid"][1]])
for r in range(2, ws2.max_row+1):
    for c in range(1, 6):
        ws2.cell(r, c).border = border; ws2.cell(r, c).alignment = center
for i, w in enumerate([12, 18, 14, 10, 10], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"

# ---- Sheet3 分类统计 ----
ws3 = wb.create_sheet("分类统计")
ws3.append(["类别", "数量", "颜色"])
for c in range(1, 4):
    cell = ws3.cell(1, c); cell.fill = head_fill; cell.font = head_font; cell.alignment = center; cell.border = border
from collections import Counter
cnt = Counter(b["cat"] for b in recs)
for k, (zh, hex_) in C.CATEGORIES.items():
    if cnt.get(k, 0) > 0:
        ws3.append([zh, cnt[k], hex_])
        r = ws3.max_row
        ws3.cell(r, 3).fill = PatternFill("solid", fgColor=hex_.lstrip("#"))
        for c in range(1, 4):
            ws3.cell(r, c).border = border; ws3.cell(r, c).alignment = center
for i, w in enumerate([16, 8, 12], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

out = os.path.join(os.path.dirname(__file__), "..", "建筑点位台账.xlsx")
wb.save(out)
print("saved", out)
